import os
import openpyxl
import requests
import time

filepath = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Pokemon.xlsx")
wb = openpyxl.load_workbook(filepath)
ws = wb.active

# Add headers for moves
headers = [cell.value for cell in ws[1]]
move_cols = []
for i, move_name in enumerate(["Move 1", "Move 2", "Move 3", "Move 4"]):
    if move_name in headers:
        col = headers.index(move_name) + 1
    else:
        col = ws.max_column + 1 + i if i == 0 else move_cols[-1] + 1
        ws.cell(row=1, column=col, value=move_name)
    move_cols.append(col)

print(f"Move columns: {move_cols}")

# Track already fetched IDs to avoid duplicate API calls (Mega variants share IDs)
cache = {}
total = ws.max_row - 1
errors = []

for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    row_num = row[0].row
    pokemon_id = row[0].value
    name = row[1].value

    if row_num % 50 == 0 or row_num == 2:
        print(f"Processing row {row_num}/{total + 1}: {name} (ID: {pokemon_id})")

    if pokemon_id in cache:
        moves = cache[pokemon_id]
    else:
        try:
            resp = requests.get(f"https://pokeapi.co/api/v2/pokemon/{pokemon_id}", timeout=10)
            if resp.status_code == 200:
                data = resp.json()
                # Get first 4 moves, capitalize nicely
                all_moves = data.get("moves", [])
                moves = []
                for m in all_moves[:4]:
                    move_name = m["move"]["name"].replace("-", " ").title()
                    moves.append(move_name)
                cache[pokemon_id] = moves
            else:
                print(f"  Warning: API returned {resp.status_code} for ID {pokemon_id} ({name})")
                moves = []
                errors.append(name)
        except Exception as e:
            print(f"  Error fetching ID {pokemon_id} ({name}): {e}")
            moves = []
            errors.append(name)

        # Small delay to be nice to the API
        time.sleep(0.1)

    # Write moves to cells
    for i, col in enumerate(move_cols):
        if i < len(moves):
            ws.cell(row=row_num, column=col, value=moves[i])
        else:
            ws.cell(row=row_num, column=col, value=None)

wb.save(filepath)
print(f"\nDone! Processed {total} Pokemon.")
if errors:
    print(f"Errors: {errors}")
else:
    print("No errors!")
