import os
import openpyxl
import requests
import time

filepath = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Pokemon.xlsx")
wb = openpyxl.load_workbook(filepath)
ws = wb.active

headers = [cell.value for cell in ws[1]]
print("Current headers:", headers)

# Find move columns
move_col_indices = []
for m in ["Move 1", "Move 2", "Move 3", "Move 4"]:
    if m in headers:
        move_col_indices.append(headers.index(m) + 1)

print(f"Move columns at: {move_col_indices}")

# Add Move Type columns
type_col_indices = []
for i, name in enumerate(["Move 1 Type", "Move 2 Type", "Move 3 Type", "Move 4 Type"]):
    if name in headers:
        col = headers.index(name) + 1
    else:
        col = ws.max_column + 1
        ws.cell(row=1, column=col, value=name)
    type_col_indices.append(col)
    # Update headers list
    headers = [cell.value for cell in ws[1]]

print(f"Move Type columns at: {type_col_indices}")

# Cache for move types (move_name -> type)
move_type_cache = {}
errors = []

def get_move_type(move_name):
    if not move_name:
        return None
    if move_name in move_type_cache:
        return move_type_cache[move_name]

    # Convert "Fire Punch" -> "fire-punch" for API
    api_name = move_name.lower().replace(" ", "-")
    try:
        resp = requests.get(f"https://pokeapi.co/api/v2/move/{api_name}", timeout=10)
        if resp.status_code == 200:
            data = resp.json()
            move_type = data["type"]["name"].title()
            move_type_cache[move_name] = move_type
            return move_type
        else:
            print(f"  Warning: API returned {resp.status_code} for move '{move_name}'")
            errors.append(move_name)
            move_type_cache[move_name] = None
            return None
    except Exception as e:
        print(f"  Error fetching move '{move_name}': {e}")
        errors.append(move_name)
        move_type_cache[move_name] = None
        return None

total = ws.max_row - 1
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    row_num = row[0].row
    name = row[1].value

    if row_num % 100 == 0 or row_num == 2:
        print(f"Processing row {row_num}/{total + 1}: {name}")

    for i, move_col in enumerate(move_col_indices):
        move_name = ws.cell(row=row_num, column=move_col).value
        move_type = get_move_type(move_name)
        ws.cell(row=row_num, column=type_col_indices[i], value=move_type)
        if move_name and move_name not in move_type_cache:
            time.sleep(0.1)

wb.save(filepath)
print(f"\nDone! Processed {total} Pokemon.")
print(f"Unique moves found: {len(move_type_cache)}")
if errors:
    print(f"Errors ({len(errors)}): {set(errors)}")
else:
    print("No errors!")
