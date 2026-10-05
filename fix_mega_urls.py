import os
import openpyxl

# Mapping of Excel Mega names to PokeAPI Mega IDs
MEGA_ID_MAP = {
    "VenusaurMega Venusaur": 10033,
    "CharizardMega Charizard X": 10034,
    "CharizardMega Charizard Y": 10035,
    "BlastoiseMega Blastoise": 10036,
    "AlakazamMega Alakazam": 10037,
    "GengarMega Gengar": 10038,
    "KangaskhanMega Kangaskhan": 10039,
    "PinsirMega Pinsir": 10040,
    "GyaradosMega Gyarados": 10041,
    "AerodactylMega Aerodactyl": 10042,
    "MewtwoMega Mewtwo X": 10043,
    "MewtwoMega Mewtwo Y": 10044,
    "AmpharosMega Ampharos": 10045,
    "ScizorMega Scizor": 10046,
    "HeracrossMega Heracross": 10047,
    "HoundoomMega Houndoom": 10048,
    "TyranitarMega Tyranitar": 10049,
    "BlazikenMega Blaziken": 10050,
    "GardevoirMega Gardevoir": 10051,
    "MawileMega Mawile": 10052,
    "AggronMega Aggron": 10053,
    "MedichamMega Medicham": 10054,
    "ManectricMega Manectric": 10055,
    "BanetteMega Banette": 10056,
    "AbsolMega Absol": 10057,
    "GarchompMega Garchomp": 10058,
    "LucarioMega Lucario": 10059,
    "AbomasnowMega Abomasnow": 10060,
    "LatiasMega Latias": 10062,
    "LatiosMega Latios": 10063,
    "SwampertMega Swampert": 10064,
    "SceptileMega Sceptile": 10065,
    "SableyeMega Sableye": 10066,
    "AltariaMega Altaria": 10067,
    "GalladeMega Gallade": 10068,
    "AudinoMega Audino": 10069,
    "SharpedoMega Sharpedo": 10070,
    "SlowbroMega Slowbro": 10071,
    "SteelixMega Steelix": 10072,
    "PidgeotMega Pidgeot": 10073,
    "GlalieMega Glalie": 10074,
    "DiancieMega Diancie": 10075,
    "MetagrossMega Metagross": 10076,
    "RayquazaMega Rayquaza": 10079,
    "CameruptMega Camerupt": 10087,
    "LopunnyMega Lopunny": 10088,
    "SalamenceMega Salamence": 10089,
    "BeedrillMega Beedrill": 10090,
}

BASE_URL = "https://raw.githubusercontent.com/PokeAPI/sprites/master/sprites/pokemon/other/official-artwork/"

filepath = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Pokemon.xlsx")
wb = openpyxl.load_workbook(filepath)
ws = wb.active

# Find or create ImageUrl column
header_row = [cell.value for cell in ws[1]]
if "ImageUrl" in header_row:
    img_col = header_row.index("ImageUrl") + 1
    print(f"Found existing ImageUrl column at column {img_col}")
else:
    img_col = ws.max_column + 1
    ws.cell(row=1, column=img_col, value="ImageUrl")
    print(f"Created ImageUrl column at column {img_col}")

updated = 0
created = 0
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    pokemon_id = row[0].value   # Column A: #
    name = row[1].value         # Column B: Name
    row_num = row[0].row

    if name in MEGA_ID_MAP:
        url = f"{BASE_URL}{MEGA_ID_MAP[name]}.png"
        updated += 1
    else:
        url = f"{BASE_URL}{pokemon_id}.png"
        created += 1

    ws.cell(row=row_num, column=img_col, value=url)

wb.save(filepath)
print(f"\nDone! Updated {updated} Mega URLs, {created} regular URLs.")
print(f"Total rows processed: {updated + created}")
